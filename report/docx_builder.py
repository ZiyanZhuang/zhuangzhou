"""
DocxBuilder — Scientific Word Document Builder
================================================
Wraps python-docx with publication-ready defaults.

Usage:
    doc = DocxBuilder('Mutation Analysis Report')
    doc.heading('Introduction', level=1)
    doc.paragraph('CYP3A4 is a key drug-metabolizing enzyme...')
    doc.table([['Mutation', 'Score'], ['V21L', '0.85']])
    doc.figure('path/to/figure.png', caption='Figure 1. Active site.')
    doc.save('report.docx')
"""

import logging
import os
from typing import List, Optional, Union

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

from . import config as C

logger = logging.getLogger(__name__)


class DocxBuilder:
    """Fluent builder for scientific Word documents."""

    def __init__(self, title: str = '', author: str = 'Zhuangzhou'):
        self.doc = Document()
        self._title = title
        self._author = author
        self._figure_count = 0
        self._table_count = 0
        self._setup_styles()
        self._setup_page()

    # ── Setup ──

    def _setup_page(self):
        """Configure A4 page with standard margins."""
        section = self.doc.sections[0]
        section.page_width = C.PAGE_WIDTH
        section.page_height = C.PAGE_HEIGHT
        section.top_margin = C.MARGIN_TOP
        section.bottom_margin = C.MARGIN_BOTTOM
        section.left_margin = C.MARGIN_LEFT
        section.right_margin = C.MARGIN_RIGHT

    def _setup_styles(self):
        """Configure document styles for scientific publications."""
        style = self.doc.styles['Normal']
        font = style.font
        font.name = C.FONT_BODY
        font.size = C.SIZE_BODY
        pf = style.paragraph_format
        pf.space_before = C.PARA_SPACING_BEFORE
        pf.space_after = C.PARA_SPACING_AFTER
        pf.line_spacing = C.LINE_SPACING

        # Set CJK font for Chinese text
        rpr = style.element.get_or_add_rPr()
        ea_font = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="{C.FONT_CN}"/>')
        rpr.append(ea_font)

    # ── Content Methods ──

    def title_page(self, title: str = None, subtitle: str = '',
                   author: str = None, institution: str = '',
                   date: str = None):
        """Add a title page."""
        import datetime
        title = title or self._title
        author = author or self._author
        date = date or datetime.date.today().strftime('%Y-%m-%d')

        # Vertical spacing
        for _ in range(4):
            self.doc.add_paragraph('')

        # Title
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = C.SIZE_TITLE
        run.font.name = C.FONT_TITLE
        run.font.color.rgb = C.COLOR_PRIMARY

        # Subtitle
        if subtitle:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(subtitle)
            run.font.size = C.SIZE_HEADING2
            run.font.color.rgb = C.COLOR_ACCENT

        # Spacer
        self.doc.add_paragraph('')

        # Author
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(author)
        run.font.size = C.SIZE_HEADING3

        # Institution
        if institution:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(institution)
            run.font.size = C.SIZE_BODY
            run.font.color.rgb = C.COLOR_CAPTION

        # Date
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(date)
        run.font.size = C.SIZE_BODY
        run.font.color.rgb = C.COLOR_CAPTION

        # Page break after title
        self.doc.add_page_break()
        return self

    def heading(self, text: str, level: int = 1):
        """Add a heading (level 1-4)."""
        h = self.doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.name = C.FONT_TITLE
            run.font.color.rgb = C.COLOR_PRIMARY
        return self

    def paragraph(self, text: str, bold: bool = False, italic: bool = False,
                  alignment: int = None, font_size: Pt = None):
        """Add a body paragraph."""
        p = self.doc.add_paragraph()
        if alignment is not None:
            p.alignment = alignment
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if font_size:
            run.font.size = font_size
        return self

    def rich_paragraph(self, segments: List[dict]):
        """Add a paragraph with mixed formatting.

        segments: [{'text': '...', 'bold': True, 'color': RGBColor(...)}, ...]
        """
        p = self.doc.add_paragraph()
        for seg in segments:
            run = p.add_run(seg.get('text', ''))
            run.bold = seg.get('bold', False)
            run.italic = seg.get('italic', False)
            if 'color' in seg:
                run.font.color.rgb = seg['color']
            if 'size' in seg:
                run.font.size = seg['size']
            if 'font' in seg:
                run.font.name = seg['font']
        return self

    def bullet_list(self, items: List[str], level: int = 0):
        """Add a bulleted list."""
        for item in items:
            p = self.doc.add_paragraph(item, style='List Bullet')
            p.paragraph_format.left_indent = Cm(1.27 * (level + 1))
        return self

    def numbered_list(self, items: List[str]):
        """Add a numbered list."""
        for item in items:
            self.doc.add_paragraph(item, style='List Number')
        return self

    def table(self, data: List[List[str]], header: bool = True,
              caption: str = None, col_widths: List[float] = None):
        """Add a formatted table.

        Args:
            data: 2D list of strings, first row is header if header=True
            header: Treat first row as header with special formatting
            caption: Optional table caption (auto-numbered)
            col_widths: Column widths in cm (optional)
        """
        if not data or not data[0]:
            return self

        n_rows = len(data)
        n_cols = len(data[0])
        tbl = self.doc.add_table(rows=n_rows, cols=n_cols, style='Table Grid')
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Set column widths
        if col_widths:
            for i, width in enumerate(col_widths):
                if i < n_cols:
                    for row in tbl.rows:
                        row.cells[i].width = Cm(width)

        for r, row_data in enumerate(data):
            for c, cell_text in enumerate(row_data):
                if c >= n_cols:
                    break
                cell = tbl.cell(r, c)
                cell.text = str(cell_text) if cell_text is not None else ''

                # Style text
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = C.SIZE_TABLE
                        run.font.name = C.FONT_BODY

                # Header row styling
                if r == 0 and header:
                    shading = parse_xml(
                        f'<w:shd {nsdecls("w")} '
                        f'w:fill="{C.COLOR_HEADER_BG}" w:val="clear"/>'
                    )
                    cell._tc.get_or_add_tcPr().append(shading)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
                            run.font.color.rgb = C.COLOR_HEADER_FG

                # Alternating row shading
                elif r > 0 and r % 2 == 0:
                    shading = parse_xml(
                        f'<w:shd {nsdecls("w")} '
                        f'w:fill="{C.COLOR_ALT_ROW}" w:val="clear"/>'
                    )
                    cell._tc.get_or_add_tcPr().append(shading)

        # Caption
        if caption:
            self._table_count += 1
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f'Table {self._table_count}. {caption}')
            run.font.size = C.SIZE_CAPTION
            run.font.color.rgb = C.COLOR_CAPTION
            run.bold = True

        return self

    def figure(self, image_path: str, caption: str = None,
               width: Cm = None, alignment: int = WD_ALIGN_PARAGRAPH.CENTER):
        """Add a figure with optional caption.

        Args:
            image_path: Path to image file (PNG, JPG, etc.)
            caption: Figure caption (auto-numbered)
            width: Image width (default: config FIGURE_WIDTH)
        """
        if not os.path.exists(image_path):
            logger.warning("Figure not found: %s", image_path)
            self.paragraph(f'[Figure missing: {os.path.basename(image_path)}]', italic=True)
            return self

        width = width or C.FIGURE_WIDTH
        p = self.doc.add_paragraph()
        p.alignment = alignment
        run = p.add_run()
        run.add_picture(image_path, width=width)

        if caption:
            self._figure_count += 1
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f'Figure {self._figure_count}. {caption}')
            run.font.size = C.SIZE_CAPTION
            run.font.color.rgb = C.COLOR_CAPTION
            run.italic = True

        return self

    def figures_side_by_side(self, images: List[dict]):
        """Add figures side by side in a table layout.

        images: [{'path': '...', 'caption': '...'}, ...]
        """
        n = len(images)
        if n == 0:
            return self

        tbl = self.doc.add_table(rows=2, cols=n)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Remove borders
        for row in tbl.rows:
            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    '<w:top w:val="none" w:sz="0"/>'
                    '<w:left w:val="none" w:sz="0"/>'
                    '<w:bottom w:val="none" w:sz="0"/>'
                    '<w:right w:val="none" w:sz="0"/>'
                    '</w:tcBorders>'
                )
                tc_pr.append(borders)

        width_per = C.FIGURE_WIDTH_FULL / n

        for i, img in enumerate(images):
            # Image cell
            cell = tbl.cell(0, i)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(img.get('path', '')):
                run = p.add_run()
                run.add_picture(img['path'], width=Cm(width_per.cm - 0.5))
            else:
                p.add_run(f'[Missing: {os.path.basename(img.get("path", "?"))}]')

            # Caption cell
            if img.get('caption'):
                self._figure_count += 1
                cap_cell = tbl.cell(1, i)
                p = cap_cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(f'Figure {self._figure_count}. {img["caption"]}')
                run.font.size = C.SIZE_CAPTION
                run.font.color.rgb = C.COLOR_CAPTION
                run.italic = True

        return self

    def code_block(self, code: str, language: str = ''):
        """Add a code block with monospace font."""
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        run = p.add_run(code)
        run.font.name = C.FONT_CODE
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        # Light background via shading
        shading = parse_xml(
            f'<w:shd {nsdecls("w")} w:fill="F5F5F5" w:val="clear"/>'
        )
        p._element.get_or_add_pPr().append(shading)
        return self

    def key_value_block(self, data: dict, title: str = None):
        """Add a key-value summary block (e.g., for metadata)."""
        if title:
            self.heading(title, level=3)

        for key, value in data.items():
            p = self.doc.add_paragraph()
            run_key = p.add_run(f'{key}: ')
            run_key.bold = True
            run_key.font.size = C.SIZE_BODY
            run_val = p.add_run(str(value))
            run_val.font.size = C.SIZE_BODY

        return self

    def horizontal_rule(self):
        """Add a horizontal line separator."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run('─' * 60)
        run.font.size = Pt(6)
        run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        return self

    def page_break(self):
        """Add a page break."""
        self.doc.add_page_break()
        return self

    def spacer(self, n: int = 1):
        """Add empty paragraphs as vertical space."""
        for _ in range(n):
            self.doc.add_paragraph('')
        return self

    # ── Specialized Scientific Content ──

    def three_line_table(self, data: List[List[str]], caption: str = None,
                         col_widths: List[float] = None, note: str = None):
        """Add a publication-standard three-line table (no internal vertical/horizontal borders).

        Structure:
          ═══════════════════  (1.5 pt top line)
          Header row (bold)
          ───────────────────  (0.75 pt header separator)
          Data rows
          ═══════════════════  (1.5 pt bottom line)

        Args:
            data:       2D list; data[0] is the header row.
            caption:    Table caption, auto-numbered.
            col_widths: Column widths in cm (optional).
            note:       Optional footnote below the table (e.g., 'a kcal/mol = ...').
        """
        if not data or not data[0]:
            return self

        n_rows = len(data)
        n_cols = len(data[0])

        # --- Caption ABOVE table (Nature/Cell style) ---
        if caption:
            self._table_count += 1
            p = self.doc.add_paragraph()
            run_label = p.add_run(f'Table {self._table_count}. ')
            run_label.bold = True
            run_label.font.size = C.SIZE_TABLE
            run_label.font.name = C.FONT_BODY
            run_body = p.add_run(caption)
            run_body.font.size = C.SIZE_TABLE
            run_body.font.name = C.FONT_BODY

        tbl = self.doc.add_table(rows=n_rows, cols=n_cols)
        tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Helper: build border XML element
        def _border_el(side: str, size: int, color: str = '000000') -> str:
            return (f'<w:{side} {nsdecls("w")} w:val="single" '
                    f'w:sz="{size}" w:space="0" w:color="{color}"/>')

        def _no_border(side: str) -> str:
            return f'<w:{side} {nsdecls("w")} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'

        # Remove table-level borders entirely (use tblPr property — CT_Tbl API)
        tbl_pr = tbl._tbl.tblPr
        tbl_borders_xml = (
            f'<w:tblBorders {nsdecls("w")}>'
            + _no_border('top') + _no_border('left') + _no_border('bottom')
            + _no_border('right') + _no_border('insideH') + _no_border('insideV')
            + '</w:tblBorders>'
        )
        existing = tbl_pr.find(qn('w:tblBorders'))
        if existing is not None:
            tbl_pr.remove(existing)
        tbl_pr.append(parse_xml(tbl_borders_xml))

        # Fill cells and apply per-cell borders
        for r_idx, row_data in enumerate(data):
            for c_idx, cell_text in enumerate(row_data):
                if c_idx >= n_cols:
                    break
                cell = tbl.cell(r_idx, c_idx)
                cell.text = str(cell_text) if cell_text is not None else ''

                # Text formatting
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = C.SIZE_TABLE
                        run.font.name = C.FONT_BODY
                        if r_idx == 0:
                            run.bold = True

                # Border: determine which edges to draw
                is_first_row = r_idx == 0
                is_last_row = r_idx == n_rows - 1

                tc_pr = cell._tc.get_or_add_tcPr()
                existing_b = tc_pr.find(qn('w:tcBorders'))
                if existing_b is not None:
                    tc_pr.remove(existing_b)

                # Build cell border XML
                top_b = _border_el('top', 12) if is_first_row else _no_border('top')
                bot_b = (_border_el('bottom', 12) if is_last_row
                         else (_border_el('bottom', 6) if is_first_row
                               else _no_border('bottom')))
                cell_borders_xml = (
                    f'<w:tcBorders {nsdecls("w")}>'
                    + top_b
                    + _no_border('left')
                    + bot_b
                    + _no_border('right')
                    + '</w:tcBorders>'
                )
                tc_pr.append(parse_xml(cell_borders_xml))

        # Set column widths
        if col_widths:
            for i, width in enumerate(col_widths):
                if i < n_cols:
                    for row in tbl.rows:
                        row.cells[i].width = Cm(width)

        # Optional footnote
        if note:
            p = self.doc.add_paragraph()
            run = p.add_run(note)
            run.font.size = Pt(8)
            run.italic = True
            run.font.color.rgb = C.COLOR_CAPTION

        return self

    def affinity_table(self, affinities: List[dict], caption: str = None):
        """Add a docking affinity results table.

        affinities: [{'rank': 1, 'affinity': -8.5, 'rmsd_lb': 0.0, 'rmsd_ub': 0.0}, ...]
        """
        header = ['Rank', 'Affinity (kcal/mol)', 'RMSD l.b. (A)', 'RMSD u.b. (A)']
        rows = [header]
        for a in affinities:
            rows.append([
                str(a.get('rank', a.get('mode', '?'))),
                f'{a.get("affinity", a.get("aff", 0)):.2f}',
                f'{a.get("rmsd_lb", 0):.3f}',
                f'{a.get("rmsd_ub", 0):.3f}',
            ])
        self.table(rows, caption=caption or 'Docking affinity scores')
        return self

    def mutation_table(self, mutations: List[dict], caption: str = None):
        """Add a mutation analysis summary table.

        mutations: [{'mutation': 'V21L', 'score': 0.85, 'role': 'active_site',
                     'recommendation': 'Promising'}, ...]
        """
        header = ['Mutation', 'Score', 'Structural Role', 'Recommendation']
        rows = [header]
        for m in mutations:
            rows.append([
                str(m.get('mutation', '?')),
                f'{m.get("score", m.get("chemical_score", 0)):.2f}',
                str(m.get('role', m.get('classification', '?'))),
                str(m.get('recommendation', '?')),
            ])
        self.table(rows, caption=caption or 'Mutation analysis results')
        return self

    def interaction_table(self, interactions: List[dict], caption: str = None):
        """Add a PLIP interaction summary table."""
        header = ['Type', 'Residue', 'Distance (A)', 'Details']
        rows = [header]
        for ix in interactions:
            rows.append([
                str(ix.get('type', '?')),
                str(ix.get('residue', ix.get('resnr', '?'))),
                f'{ix.get("distance", ix.get("dist", 0)):.2f}',
                str(ix.get('details', ix.get('info', ''))),
            ])
        self.table(rows, caption=caption or 'Protein-ligand interactions')
        return self

    def energy_summary(self, energies: dict):
        """Add an energy summary block (for QM/MM or MD results)."""
        self.heading('Energy Summary', level=3)
        items = []
        for key, val in energies.items():
            if isinstance(val, float):
                items.append(f'{key}: {val:.4f} kcal/mol')
            else:
                items.append(f'{key}: {val}')
        self.bullet_list(items)
        return self

    # ── Output ──

    def save(self, output_path: str) -> str:
        """Save the document and return the absolute path."""
        output_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        self.doc.save(output_path)
        logger.info("Report saved: %s", output_path)
        return output_path

    def __repr__(self):
        return (f'DocxBuilder(title="{self._title}", '
                f'figures={self._figure_count}, tables={self._table_count})')
